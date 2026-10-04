#!/usr/bin/env python3
"""Read addresses from, and write Zillow rent estimates into, the MKE screening workbook.

Subcommands:
  list   Print JSON for every row that still needs a Zillow rent estimate.
  write  Write one estimate (and optional source URL) into a row.

The target column is found by header text ("Zillow Rent Zestimate") unless
--column is given, so it works on Sheet1 (column AB) and Favorties (column AC)
without hard-coding letters. The source-URL column ("Zillow Source") is used
when the sheet has one.
"""
import argparse
import json
import sys

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

DEFAULT_STATE = "WI"
TARGET_HEADER = "zillow rent zestimate"
SOURCE_HEADER = "zillow source"
# Formula-driven columns we refuse to overwrite unless --force is passed.
PROTECTED_HEADER_HINTS = ("req rent", "implied cap", "rent gap", "screen", "units")


def load(path):
    keep_vba = path.lower().endswith(".xlsm")
    return openpyxl.load_workbook(path, keep_vba=keep_vba)


def find_header_row(ws, max_scan=15):
    """Return the first row whose column A says 'address'."""
    for r in range(1, max_scan + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip().lower() == "address":
            return r
    raise SystemExit(f"Could not find an 'address' header in column A of '{ws.title}'.")


def header_map(ws, header_row):
    out = {}
    for c in ws[header_row]:
        if isinstance(c.value, str):
            out[c.value.strip().lower()] = c.column
    return out


def find_col(headers, needle):
    for text, col in headers.items():
        if needle in text:
            return col
    return None


def resolve_columns(ws, header_row, column=None, force=False):
    headers = header_map(ws, header_row)
    if column:
        target = column_index_from_string(column.upper())
        hdr = (ws.cell(header_row, target).value or "")
        if not force and any(h in str(hdr).lower() for h in PROTECTED_HEADER_HINTS):
            raise SystemExit(
                f"Column {column} on '{ws.title}' is '{hdr}', a formula column. "
                "Pass --force to overwrite it anyway."
            )
    else:
        target = find_col(headers, TARGET_HEADER)
        if target is None:
            raise SystemExit(
                f"No 'Zillow Rent Zestimate' header on '{ws.title}'. Pass --column (e.g. --column AB)."
            )
    source = find_col(headers, SOURCE_HEADER)
    city = find_col(headers, "city")
    return target, source, city


def full_address(street, city, state=DEFAULT_STATE):
    """'2025 South 76th Street, Unit 2027' + 'West Allis' -> '2025 South 76th Street, Unit 2027, West Allis, WI'.

    Any unit already in the street cell is kept as-is.
    """
    parts = [str(street).strip()]
    if city:
        parts.append(str(city).strip())
    parts.append(state)
    return ", ".join(parts)


def iter_rows(ws, header_row, target, city_col, overwrite=False, rows=None):
    for r in range(header_row + 1, ws.max_row + 1):
        if rows and r not in rows:
            continue
        street = ws.cell(r, 1).value
        if not street or not str(street).strip():
            continue
        current = ws.cell(r, target).value
        if current not in (None, "") and not overwrite:
            continue
        city = ws.cell(r, city_col).value if city_col else None
        yield {
            "row": r,
            "street": str(street).strip(),
            "city": city,
            "address": full_address(street, city),
            "current": current,
        }


def parse_rows(spec):
    if not spec:
        return None
    rows = set()
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            rows.update(range(int(a), int(b) + 1))
        elif part:
            rows.add(int(part))
    return rows


def write_value(ws, row, target, source_col, value, source_url=None):
    cell = ws.cell(row, target)
    cell.value = value
    if isinstance(value, (int, float)):
        cell.number_format = '"$"#,##0'
    if source_col and source_url:
        ws.cell(row, source_col).value = source_url


def cmd_list(args):
    wb = load(args.workbook)
    ws = wb[args.sheet]
    hr = find_header_row(ws)
    target, source, city = resolve_columns(ws, hr, args.column, args.force)
    items = list(iter_rows(ws, hr, target, city, args.overwrite, parse_rows(args.rows)))
    json.dump(
        {
            "sheet": ws.title,
            "target_column": get_column_letter(target),
            "source_column": get_column_letter(source) if source else None,
            "count": len(items),
            "rows": items,
        },
        sys.stdout,
        indent=2,
        default=str,
    )
    print()


def cmd_write(args):
    wb = load(args.workbook)
    ws = wb[args.sheet]
    hr = find_header_row(ws)
    target, source, _ = resolve_columns(ws, hr, args.column, args.force)
    value = args.value
    try:
        value = int(round(float(str(value).replace("$", "").replace(",", ""))))
    except ValueError:
        pass  # keep text such as "N/A"
    write_value(ws, args.row, target, source, value, args.source)
    wb.save(args.out or args.workbook)
    print(f"{ws.title}!{get_column_letter(target)}{args.row} = {value}")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("workbook")
    p.add_argument("--sheet", default="Sheet1")
    p.add_argument("--column", help="Target column letter; default finds the 'Zillow Rent Zestimate' header")
    p.add_argument("--force", action="store_true", help="Allow writing into a formula column")
    sub = p.add_subparsers(dest="cmd", required=True)

    pl = sub.add_parser("list")
    pl.add_argument("--overwrite", action="store_true", help="Include rows that already have a value")
    pl.add_argument("--rows", help="Only these sheet rows, e.g. 5-20,31")
    pl.set_defaults(func=cmd_list)

    pw = sub.add_parser("write")
    pw.add_argument("--row", type=int, required=True)
    pw.add_argument("--value", required=True, help="Monthly rent, e.g. 1650 or $1,650; or text like N/A")
    pw.add_argument("--source", help="URL to record in the 'Zillow Source' column, if the sheet has one")
    pw.add_argument("--out", help="Save to a different file instead of in place")
    pw.set_defaults(func=cmd_write)

    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
